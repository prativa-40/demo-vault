from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs' / 'Hearth_Product_Description_and_Installation_Manual.docx'
IMAGE = ROOT / 'public' / 'assets' / 'sanduk-hero.png'

INK = '14334A'; TEAL = '2E6F6A'; GOLD = 'B77B37'; PALE = 'EAF3F1'; PALE_BLUE = 'E8EEF5'; GREY = '55636D'; LIGHT = 'F5F7F7'

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), fill); tc_pr.append(shade)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); mar = tc_pr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tc_pr.append(mar)
    for side, value in [('top',top), ('start',start), ('bottom',bottom), ('end',end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None: node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def fix_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None: tbl_w = OxmlElement('w:tblW'); tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths))); tbl_w.set(qn('w:type'), 'dxa')
    ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa'); tbl_pr.append(ind)
    grid = table._tbl.tblGrid
    for i, width in enumerate(widths):
        grid.gridCol_lst[i].set(qn('w:w'), str(width))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in('w:tcW')
            tc_w.set(qn('w:w'), str(widths[i])); tc_w.set(qn('w:type'), 'dxa')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

def font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = 'Calibri'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold; run.italic = italic

def paragraph(doc, text='', size=11, color=INK, bold=False, italic=False, after=6, before=0, align=None, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.25
    if align is not None: p.alignment = align
    if text: font(p.add_run(text), size, color, bold, italic)
    return p

def heading(doc, text, level=1):
    sizes = {1: 16, 2: 13, 3: 12}; colors = {1: TEAL, 2: TEAL, 3: INK}
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt({1:18,2:14,3:10}[level]); p.paragraph_format.space_after = Pt({1:10,2:7,3:5}[level]); p.paragraph_format.keep_with_next = True
    font(p.add_run(text), sizes[level], colors[level], True)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.28); p.paragraph_format.first_line_indent = Inches(-0.20); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.25
    font(p.add_run('• '), 10.5, TEAL, True); font(p.add_run(text), 10.5, INK); return p

def step(doc, text):
    doc._hearth_step = getattr(doc, '_hearth_step', 0) + 1
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.28); p.paragraph_format.first_line_indent = Inches(-0.28); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.25
    font(p.add_run(f'Step {doc._hearth_step}. '), 10.5, TEAL, True); font(p.add_run(text), 10.5, INK); return p

def reset_steps(doc):
    doc._hearth_step = 0

def label_value(doc, label, value):
    p = paragraph(doc, after=3); font(p.add_run(label + ': '), 10.5, INK, True); font(p.add_run(value), 10.5, INK)

def set_header_footer(section):
    header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0); font(header.add_run('HEARTH  |  PRODUCT & INSTALLATION GUIDE'), 8.5, GREY, True)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0); font(footer.add_run('Private household records, kept close to home'), 8.5, GREY)

doc = Document(); section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(1); section.left_margin = section.right_margin = Inches(1)
section.header_distance = Inches(0.492); section.footer_distance = Inches(0.492); set_header_footer(section)

normal = doc.styles['Normal']; normal.font.name='Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); normal.font.size=Pt(11)

# Cover
paragraph(doc, 'HEARTH', 12, GOLD, True, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
p = paragraph(doc, 'Your private household vault', 29, INK, True, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
paragraph(doc, 'Product description and installation manual', 14, TEAL, False, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
if IMAGE.exists():
    doc.add_picture(str(IMAGE), width=Inches(3.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph(doc, 'A local-first system for family records, documents, reminders, and private AI assistance.', 11.5, GREY, False, after=16, before=14, align=WD_ALIGN_PARAGRAPH.CENTER)
table = doc.add_table(rows=3, cols=2); fix_table(table, [2100, 7260])
for row, (left, right) in zip(table.rows, [('Audience', 'Households, home-server operators, and implementation partners'), ('Deployment model', 'Private home network with encrypted storage and locally hosted AI'), ('Version', 'Product guide — July 2026')]):
    set_cell_shading(row.cells[0], PALE); font(row.cells[0].paragraphs[0].add_run(left), 10, INK, True); font(row.cells[1].paragraphs[0].add_run(right), 10, INK)

doc.add_page_break()

heading(doc, '1. Product overview', 1)
paragraph(doc, 'Hearth is a private household information system. It brings together family profiles, identity and education records, health information, finance, property, vehicles, and time-sensitive reminders in one encrypted home vault. The browser interface is used to add, review, retrieve, and update records; the home server is the system of record.', after=8)
heading(doc, 'What Hearth solves', 2)
for item in ['Keeps essential household records organised by family member and category.', 'Stores original scans separately from metadata, with encryption at rest for both.', 'Reads uploaded PDFs and images locally to suggest document type, name, document number, date of birth, address, and other metadata for review.', 'Supports private search and question answering over household records when the local RAG service is configured.', 'Creates a foundation for expiry, birthday, renewal, warranty, insurance, and service reminders.']:
    bullet(doc, item)

heading(doc, 'Core product areas', 2)
table = doc.add_table(rows=1, cols=2); fix_table(table, [2350, 7010]); set_cell_shading(table.rows[0].cells[0], PALE_BLUE); set_cell_shading(table.rows[0].cells[1], PALE_BLUE)
font(table.rows[0].cells[0].paragraphs[0].add_run('Area'), 10, INK, True); font(table.rows[0].cells[1].paragraphs[0].add_run('Typical records'), 10, INK, True)
for area, details in [('Family Members', 'Name, relationship, blood group, gender, household role and linked records.'), ('Identity & Documents', 'Aadhaar, passport, voter ID, PAN, driving licence, birth and marriage certificates, scanned copies and extracted metadata.'), ('Academic & Achievements', 'Class 10, Class 12, graduation, postgraduate, PhD, sport and extracurricular records.'), ('Finance', 'Banking, cards, UPI, fixed deposits, PPF, EPF, NPS, demat, insurance, loans and import templates.'), ('Health', 'Allergies, medication, reports, vaccinations and doctor contacts.'), ('House & Property', 'Ownership, deeds, utility bills, subscriptions, appliances, warranties and pest-control records.'), ('Vehicles & Trackers', 'RC, insurance, PUC, challans, maintenance, due dates and renewal tracking.')]:
    cells=table.add_row().cells; font(cells[0].paragraphs[0].add_run(area),10,INK,True); font(cells[1].paragraphs[0].add_run(details),10,INK)

heading(doc, 'Privacy and security design', 2)
paragraph(doc, 'Hearth is designed for a private home deployment. Records and original documents are encrypted at rest with AES-256-GCM. The web server does not expose the encrypted data directory as a public static folder; documents are delivered through private API routes.', after=4)
note = doc.add_table(rows=1, cols=1); fix_table(note, [9360]); set_cell_shading(note.cell(0,0), 'FFF4DE'); font(note.cell(0,0).paragraphs[0].add_run('Deployment note. The current prototype provides encrypted-at-rest storage. Before allowing access outside the trusted home network, add TLS, authentication, server-enforced permissions, audit logging, backups, and managed encryption keys.'), 10, '684800', False)

heading(doc, '2. Architecture at a glance', 1)
paragraph(doc, 'Each household device uses the web application over the home network. The API engine receives records and uploads, passes files to the ingestion service, and reads or writes the encrypted records database and document store. Local AI services run on the home server, not in a public cloud.', after=7)
table = doc.add_table(rows=1, cols=3); fix_table(table, [2450, 3450, 3460])
for cell, text in zip(table.rows[0].cells, ['Layer', 'Current implementation', 'Purpose']):
    set_cell_shading(cell, PALE_BLUE); font(cell.paragraphs[0].add_run(text), 10, INK, True)
for row in [('Web application', 'Safari / desktop browser UI', 'Family record management, document review, retrieval and role-aware screens.'), ('API engine', 'Node.js server.js', 'Private routes, upload handling, document delivery and application logic.'), ('Ingestion service', 'OCR + document-intake agent', 'Reads PDFs/images and prepares reviewable metadata.'), ('Encrypted records database', 'data/household.enc', 'Members, assets, dates, metadata, permissions and RAG index.'), ('Encrypted document store', 'data/documents/', 'Original scans and uploaded files.'), ('Local AI & RAG', 'Ollama-compatible service', 'Private document understanding, embeddings, semantic retrieval and cited answers.')]:
    cells=table.add_row().cells
    for cell, text in zip(cells,row): font(cell.paragraphs[0].add_run(text),9.5,INK)

heading(doc, '3. Installation prerequisites', 1)
label_value(doc, 'Home server', 'A Mac, Linux machine, mini PC, or dedicated Hearth appliance that stays available on the home network.')
label_value(doc, 'Node.js', 'A current Node.js runtime. The project uses only Node built-ins and starts with npm start.')
label_value(doc, 'OCR', 'macOS provides the included Vision/PDFKit OCR script. A Linux deployment requires an equivalent OCR implementation.')
label_value(doc, 'Local AI (optional but recommended)', 'Ollama-compatible local service with a chat model and an embedding model. Example: qwen3:8b plus a local embedding model.')
label_value(doc, 'Network', 'Trusted home LAN. Keep the initial prototype bound to the local/private network until production controls are added.')

heading(doc, '4. Install Hearth', 1)
paragraph(doc, 'These steps install the current local prototype on the machine that will act as the Hearth home server.', after=7)
heading(doc, 'A. Obtain the source', 2)
reset_steps(doc)
step(doc, 'Copy or clone the Hearth project folder onto the home server.')
step(doc, 'Open a terminal in the project folder.')
paragraph(doc, 'Expected project files include server.js, package.json, lib/, public/, scripts/, and docs/.', 10, GREY, italic=True, after=8)
heading(doc, 'B. Start the web application', 2)
step(doc, 'Run npm start.')
step(doc, 'Open http://127.0.0.1:3000 in Safari on the home server.')
step(doc, 'For another device on the trusted home network, configure the server/network binding deliberately before exposing access. Do not publish the prototype directly to the internet.')
heading(doc, 'C. Configure the production encryption key', 2)
paragraph(doc, 'For a production-style launch, provide a base64-encoded 32-byte key through the server environment. Do not store this secret in source control or in a shared note.', after=4)
code = doc.add_table(rows=1, cols=1); fix_table(code, [9360]); set_cell_shading(code.cell(0,0), '1F2933'); font(code.cell(0,0).paragraphs[0].add_run('export HEARTH_MASTER_KEY="<base64-encoded-32-byte-secret>"\nnpm start'), 9.5, 'FFFFFF')
paragraph(doc, 'In development, Hearth creates a local development key with owner-only permissions. Treat it as development-only.', 10, GREY, italic=True, after=8)

doc.add_page_break()
heading(doc, '5. Configure private AI and RAG', 1)
paragraph(doc, 'Hearth uses a local Ollama-compatible endpoint. No cloud-AI call is required by the product code. The chat model supports document intake and question answering; the embedding model supports semantic retrieval.', after=7)
reset_steps(doc)
step(doc, 'Install and start your local Ollama-compatible service on the Hearth server.')
step(doc, 'Download the chat model and an embedding model suited to your environment.')
step(doc, 'Set the environment variables before starting Hearth.')
code = doc.add_table(rows=1, cols=1); fix_table(code, [9360]); set_cell_shading(code.cell(0,0), '1F2933'); font(code.cell(0,0).paragraphs[0].add_run('export LOCAL_LLM_URL="http://127.0.0.1:11434"\nexport LOCAL_LLM_MODEL="qwen3:8b"\nexport LOCAL_EMBEDDING_MODEL="<your-local-embedding-model>"\nnpm start'), 9.5, 'FFFFFF')
paragraph(doc, 'After Hearth starts, use the private AI area to reindex the vault. This creates encrypted vector entries from approved records and OCR text. Reindex after important uploads or edits.', after=6)

heading(doc, '6. First-use checklist', 1)
for item in ['Open the Overview screen and create the household’s family members.', 'Assign household roles: Head of household, Homemaker, or Child. Only the head should manage permissions.', 'Add identity and academic records with scanned copies.', 'Use “Extract details” to review OCR suggestions. Confirm the member, name, DOB, number, and address before saving.', 'Add property, finance, health, vehicle, and subscription records as needed.', 'Use the document search to retrieve scans. “Open scan” opens the saved scan in a new tab.', 'Create a secure backup policy for both the encrypted records file and encrypted documents directory.']:
    bullet(doc, item)

heading(doc, '7. Document upload and review workflow', 1)
paragraph(doc, 'Hearth’s intake workflow is review-first. It does not automatically treat OCR output as final truth.', after=6)
reset_steps(doc)
for item in ['Choose a family member and document type, then select a PDF or image scan.', 'Select “Upload & extract details.” Hearth performs local OCR and, when configured, asks the local AI service to classify the scan and propose metadata.', 'Review the suggested name, date of birth, address, document number, and assigned family member.', 'Correct any uncertain value. If the DOB is not supported by visible scan evidence, leave it blank or add it manually after review.', 'Select “Save reviewed document.” The scan is encrypted and stored with the final metadata.']:
    step(doc, item)

heading(doc, '8. Operations, backup, and troubleshooting', 1)
heading(doc, 'Backups', 2)
paragraph(doc, 'Back up the encrypted records file and the encrypted document store together. A backup of only one is incomplete. Safeguard the production encryption key separately; without the matching key, encrypted data cannot be recovered.', after=7)
heading(doc, 'Common checks', 2)
table = doc.add_table(rows=1, cols=2); fix_table(table, [3100,6260]); set_cell_shading(table.rows[0].cells[0], PALE_BLUE); set_cell_shading(table.rows[0].cells[1], PALE_BLUE)
font(table.rows[0].cells[0].paragraphs[0].add_run('Situation'),10,INK,True); font(table.rows[0].cells[1].paragraphs[0].add_run('What to check'),10,INK,True)
for issue, action in [('Safari cannot open the server', 'Confirm npm start is running, then open the configured local URL. Check the chosen port and firewall/network settings.'), ('OCR cannot read a scan', 'Use a clearer, straight scan with adequate resolution. For PDFs, verify the file opens locally. Use the document’s Re-read action after improving OCR settings.'), ('Name or DOB is missing', 'Review the raw scan and OCR result. Add metadata manually if it is not readable. Do not accept a date that is not visible on the scan.'), ('Local AI is unavailable', 'Confirm LOCAL_LLM_URL, LOCAL_LLM_MODEL, and LOCAL_EMBEDDING_MODEL are exported in the same shell/service that starts Hearth. Confirm the local model service is running.'), ('Search does not reflect a new document', 'Run the private AI reindex action after the upload has been reviewed and saved.'), ('A scan will not open', 'Check that the server is running and that the encrypted file remains in data/documents/. Do not move encrypted files independently of the records database.')]:
    cells=table.add_row().cells; font(cells[0].paragraphs[0].add_run(issue),9.5,INK,True); font(cells[1].paragraphs[0].add_run(action),9.5,INK)

doc.add_page_break()
heading(doc, '9. Production readiness roadmap', 1)
paragraph(doc, 'The current implementation is a strong local prototype. Before a broader household or remote deployment, complete the following controls:', after=5)
for item in ['Server-enforced authentication, sessions, and household-role authorization.', 'TLS certificates for every device-to-server connection.', 'Audit logging for document access, exports, role changes, and destructive actions.', 'Managed keys, per-household encryption keys, secure backup and recovery drills.', 'Malware scanning, file-size limits, background ingestion jobs, and durable notification scheduling.', 'Permission-filtered RAG retrieval so each family role receives only authorised results.']:
    bullet(doc, item)

heading(doc, 'Support information', 1)
paragraph(doc, 'For a support request, provide the Hearth version, server operating system, browser, the relevant document type, whether the file is PDF or image, and the exact message shown by the application. Do not send sensitive scans or document numbers through unsecured channels.', after=0)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
